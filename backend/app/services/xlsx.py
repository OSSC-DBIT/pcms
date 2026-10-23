"""Excel import and export helpers.

Every exported cell is written as text so a name or note starting with "="
cannot become a formula. Imports reject formula cells and macro workbooks.
"""

import io
import zipfile
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.errors import validation

CREDENTIAL_HEADERS = (
    "email", "name", "role", "status", "password", "password_changed_at",
    "last_login_at", "temp_password_expires_at", "notes",
)

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_PASSWORD_FILL = PatternFill("solid", fgColor="FFF2CC")


def _as_text(value) -> str:
    if value is None:
        return ""
    return str(value)


def write_text_cell(ws, row: int, col: int, value, *, fill=None):
    cell = ws.cell(row=row, column=col, value=_as_text(value))
    # Force a string even when the value starts with "=", "+", "-" or "@".
    cell.data_type = "s"
    cell.number_format = "@"
    cell.alignment = Alignment(vertical="top", wrap_text=False)
    if fill is not None:
        cell.fill = fill
    return cell


def _style_header(ws, headers, widths):
    for i, (name, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _reject_unsafe_archive(data: bytes) -> None:
    # openpyxl drops macro and external-link parts before we can inspect the
    # loaded workbook, so check the ZIP container itself first.
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = [name.lower() for name in archive.namelist()]
    except (zipfile.BadZipFile, OSError) as exc:
        raise validation("This is not a valid .xlsx file.", {"file": str(exc)[:200]})
    if any(name.endswith("vbaproject.bin") or name.startswith("xl/macrosheets/") for name in names):
        raise validation("Macro workbooks are not accepted. Save as .xlsx.", {"file": "Macros found"})
    if any(name.startswith("xl/externallinks/") for name in names):
        raise validation("Workbooks with external links are not accepted.", {"file": "External links found"})


def read_people_workbook(data: bytes) -> list[tuple[int, list[str]]]:
    """Return (line number, cells) for the first worksheet, header included."""
    _reject_unsafe_archive(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # openpyxl raises several formats for a bad zip
        raise validation("This is not a valid .xlsx file.", {"file": str(exc)[:200]})
    try:
        if not wb.worksheets:
            raise validation("The workbook has no sheets.", {"file": "No sheets"})
        ws = wb.worksheets[0]
        out: list[tuple[int, list[str]]] = []
        for line, cells in enumerate(ws.iter_rows(values_only=False), start=1):
            values: list[str] = []
            for cell in cells:
                if getattr(cell, "data_type", None) == "f":
                    raise validation("Formulas are not accepted in the import file.",
                                     {"file": f"Formula in {cell.coordinate}"})
                values.append("" if cell.value is None else str(cell.value).strip())
            out.append((line, values))
        return out
    finally:
        wb.close()


def build_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    headers = ("email", "name", "role", "password")
    _style_header(ws, headers, (34, 28, 22, 26))
    ws.append(["", "", "", ""])

    note = wb.create_sheet("Read me")
    note.column_dimensions["A"].width = 110
    lines = [
        "PCMS people import",
        "",
        "Columns: email, name, role, password.",
        "role is one of faculty, facility_manager or management. Leave it blank for faculty.",
        "password is the temporary password Team OSSC gives the person. It must be at least 10 characters.",
        "Do not create OSSC accounts from this file. Add a super admin one at a time in People.",
        "Emails must be unique. Existing emails are skipped, never updated or reset.",
        "If any line is wrong, the whole file is rejected and nothing is imported.",
        "After import, the system downloads a credentials file. Share each row only with that person.",
        "Delete the downloaded credentials file after sharing it.",
    ]
    for i, line in enumerate(lines, start=1):
        write_text_cell(note, i, 1, line)

    return _save(wb)


def build_credentials_workbook(rows, *, generated_by: str, generated_at: datetime) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Credentials"
    _style_header(ws, CREDENTIAL_HEADERS, (34, 26, 18, 26, 24, 24, 22, 24, 48))

    for r, row in enumerate(rows, start=2):
        values = [row.get(h, "") for h in CREDENTIAL_HEADERS]
        for c, value in enumerate(values, start=1):
            fill = _PASSWORD_FILL if CREDENTIAL_HEADERS[c - 1] == "password" and value else None
            write_text_cell(ws, r, c, value, fill=fill)

    note = wb.create_sheet("Read me")
    note.column_dimensions["A"].width = 110
    local = generated_at.astimezone(ZoneInfo(get_settings().timezone)).strftime("%Y-%m-%d %H:%M IST")
    lines = [
        "PCMS credentials file - CONFIDENTIAL",
        "",
        f"Generated by: {generated_by}",
        f"Generated at: {local}",
        "",
        "This file contains temporary passwords. Share each row only with the person named in it.",
        "Do not upload it to a group chat, shared drive or WhatsApp group.",
        "Delete the file after the credentials have been shared.",
        "",
        "The password column is filled only until the person changes their temporary password.",
        "After that it is blank. Use Reset in Administration > People to issue a new one.",
        "A password chosen by an employee is never stored in a form Team OSSC can read.",
        "",
        "OSSC accounts also need a 6-digit authenticator code after the password.",
        "Status values: Not activated, Temporary password expired, Active, Suspended.",
    ]
    for i, line in enumerate(lines, start=1):
        write_text_cell(note, i, 1, line)

    return _save(wb)


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ist_stamp(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.astimezone(ZoneInfo(get_settings().timezone)).strftime("%Y-%m-%d %H:%M IST")
