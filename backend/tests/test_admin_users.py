import csv
import io
import zipfile
from datetime import timedelta

import httpx
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.db import sessionmaker
from app.main import app
from app.models import AuditEvent, User
from app.security.sessions import now
from tests.factories import API, PASSWORD, make_user, sign_in

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ADMIN = "admin@dbit.edu"
FACULTY = "faculty@dbit.edu"


async def _ossc(client):
    await make_user(ADMIN, "ossc", "OSSC Lead")
    return await sign_in(client, ADMIN)


def _csv(text: str):
    return {"file": ("people.csv", text.encode("utf-8"), "text/csv")}


def _xlsx_bytes(rows: list[list[str]], headers=("email", "name", "role", "password")) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_upload(rows: list[list[str]], headers=("email", "name", "role", "password")):
    return {"file": ("people.xlsx", _xlsx_bytes(rows, headers), XLSX)}


def _xlsx_with_member(rows: list[list[str]], member: str, *, filename="people.xlsx"):
    buf = io.BytesIO(_xlsx_bytes(rows))
    with zipfile.ZipFile(buf, "a") as archive:
        archive.writestr(member, b"not used")
    return {"file": (filename, buf.getvalue(), XLSX)}


def _rows(resp) -> list[dict]:
    wb = load_workbook(io.BytesIO(resp.content), data_only=False)
    ws = wb["Credentials"]
    it = ws.iter_rows(values_only=True)
    headers = next(it)
    return [{h: ("" if cell is None else cell) for h, cell in zip(headers, row, strict=True)}
            for row in it if any(cell is not None for cell in row)]


async def test_only_ossc_can_manage_users(client):
    for email, role in ((FACULTY, "faculty"), ("manager@dbit.edu", "facility_manager"),
                        ("management@dbit.edu", "management")):
        await make_user(email, role)
        s = await sign_in(client, email)
        assert (await s.get(f"{API}/admin/users")).status_code == 403
    assert (await s.post(f"{API}/admin/users/import", files=_csv("email,name,password\nx@dbit.edu,X,Temp-X-2345\n"))).status_code == 403
    assert (await s.post(f"{API}/admin/users/credentials/export", json={"password": PASSWORD})).status_code == 403


async def test_csv_import_returns_credentials_and_forces_change(client):
    s = await _ossc(client)
    await make_user("existing@dbit.edu", "faculty", "Existing Person")
    text = (
        "email,name,role,password\n"
        "priya@dbit.edu,Priya Nair,faculty,Temp-Person-2345\n"
        "rahul@dbit.edu,राहुल पाटील,,Temp-Person-2346\n"
        "a@dbit.edu,=HYPERLINK(\"x\"),facility_manager,Temp-Another-2345\n"
        "existing@dbit.edu,Existing Person,faculty,Temp-Other-2345\n"
        "\n"
    )
    r = await s.post(f"{API}/admin/users/import", files=_csv(text))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX)
    assert "no-store" in r.headers["cache-control"]
    assert r.headers["x-pcms-created"] == "3" and r.headers["x-pcms-existing"] == "1"
    rows = {row["email"]: row for row in _rows(r)}
    assert rows["priya@dbit.edu"]["status"] == "Not activated"
    assert rows["priya@dbit.edu"]["password"] == "Temp-Person-2345"
    assert rows["rahul@dbit.edu"]["role"] == "faculty" and rows["rahul@dbit.edu"]["name"] == "राहुल पाटील"
    assert rows["a@dbit.edu"]["name"] == '=HYPERLINK("x")'  # written as text, not a formula
    assert rows["existing@dbit.edu"]["status"] == "Already exists" and rows["existing@dbit.edu"]["password"] == ""
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Credentials"]
    name_col = [c.value for c in ws[1]].index("name") + 1
    formula_name = next(ws.cell(row=n, column=name_col) for n in range(2, ws.max_row + 1)
                        if ws.cell(row=n, column=1).value == "a@dbit.edu")
    assert formula_name.data_type == "s" and formula_name.number_format == "@"

    # The plaintext is encrypted at rest, never hashed into the audit trail.
    async with sessionmaker()() as db:
        u = await db.scalar(select(User).where(User.email == "priya@dbit.edu"))
        assert u.must_change_password and u.password_hash.startswith("$argon2id$")
        assert u.temp_password_enc and "Temp-Person-2345" not in u.temp_password_enc
        audits = (await db.scalars(select(AuditEvent).where(AuditEvent.action == "user.imported"))).all()
        assert len(audits) == 1 and "Temp-Person-2345" not in str(audits[0].after)

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as person:
        r = await person.post(f"{API}/session", json={"email": "priya@dbit.edu", "password": "Temp-Person-2345"})
        assert r.json()["next"] == "change_password"


async def test_xlsx_import_and_template(client):
    s = await _ossc(client)
    template = await s.get(f"{API}/admin/users/import/template")
    assert template.status_code == 200 and template.headers["content-type"].startswith(XLSX)
    wb = load_workbook(io.BytesIO(template.content))
    assert [c.value for c in wb["People"][1]] == ["email", "name", "role", "password"]

    r = await s.post(f"{API}/admin/users/import", files=_xlsx_upload([
        ["nina@dbit.edu", "Nina D", "management", "Temp-Person-2347"],
        ["omar@dbit.edu", "Omar S", "", "Temp-Person-2348"],
    ]))
    assert r.status_code == 200, r.text
    rows = {row["email"]: row for row in _rows(r)}
    assert rows["nina@dbit.edu"]["role"] == "management"
    assert rows["omar@dbit.edu"]["role"] == "faculty"


async def test_import_is_all_or_nothing(client):
    s = await _ossc(client)
    text = (
        "email,name,role,password\n"
        "good@dbit.edu,Good Row,faculty,Temp-Person-2349\n"
        "not-an-email,X,faculty,Temp-Bad-2345\n"
        "good@dbit.edu,Dup,faculty,Temp-Dup-2345\n"
        "ossc2@dbit.edu,Bad Role,ossc,Temp-Role-2345\n"
    )
    r = await s.post(f"{API}/admin/users/import", files=_csv(text))
    assert r.status_code == 400
    errs = r.json()["field_errors"]
    assert set(errs) == {"line 3", "line 4", "line 5"}
    async with sessionmaker()() as db:
        assert await db.scalar(select(func.count()).select_from(User)) == 1  # only the OSSC account


async def test_import_rejects_bad_files(client):
    s = await _ossc(client)
    assert (await s.post(f"{API}/admin/users/import", files=_csv("id,full_name\n1,2\n"))).status_code == 400
    assert (await s.post(f"{API}/admin/users/import", files={"file": ("x.csv", b"\xff\xfe\x00bad", "text/csv")})).status_code == 400
    assert (await s.post(f"{API}/admin/users/import", files=_csv("email,name,salary,password\na@dbit.edu,A,3,Temp-Another-2345\n"))).status_code == 400
    big = "email,name,role,password\n" + "".join(f"e{i}@dbit.edu,Person {i},faculty,Temp-Person-{i}\n" for i in range(1001))
    assert (await s.post(f"{API}/admin/users/import", files=_csv(big))).status_code == 400

    wb = Workbook()
    ws = wb.active
    ws.append(["email", "name", "role", "password"])
    ws.append(["formula@dbit.edu", "=1+1", "faculty", "Temp-Formula-2345"])
    buf = io.BytesIO()
    wb.save(buf)
    r = await s.post(f"{API}/admin/users/import", files={"file": ("f.xlsx", buf.getvalue(), XLSX)})
    assert r.status_code == 400 and "Formula" in r.text

    rows = [["macro@dbit.edu", "Macro", "faculty", "Temp-Person-2351"]]
    r = await s.post(f"{API}/admin/users/import", files=_xlsx_with_member(rows, "xl/vbaProject.bin"))
    assert r.status_code == 400 and "Macro workbooks" in r.text
    r = await s.post(f"{API}/admin/users/import",
                     files={"file": ("people.xlsm", _xlsx_bytes(rows), XLSX)})
    assert r.status_code == 400 and "Macro-enabled" in r.text
    r = await s.post(f"{API}/admin/users/import",
                     files=_xlsx_with_member(rows, "xl/externalLinks/externalLink1.xml"))
    assert r.status_code == 400 and "external links" in r.text.lower()


async def test_dry_run_creates_nothing(client):
    s = await _ossc(client)
    r = await s.post(f"{API}/admin/users/import?dryRun=true",
                     files=_csv(f"email,name,role,password\nnew@dbit.edu,Some One,faculty,Temp-Person-2350\n{ADMIN},Me,faculty,Temp-Me-2345\n"))
    assert r.status_code == 200, r.text
    assert r.json() == {"toCreate": 1, "alreadyExist": [ADMIN]}
    async with sessionmaker()() as db:
        assert await db.scalar(select(func.count()).select_from(User)) == 1


async def test_create_and_reset_user(client):
    s = await _ossc(client)
    r = await s.post(f"{API}/admin/users", json={
        "email": "r.patil@dbit.edu", "name": "R. Patil", "role": "facility_manager",
        "password": "Temp-Rpatil-2345",
    })
    assert r.status_code == 201 and "no-store" in r.headers["cache-control"]
    created = r.json()
    assert created["user"]["email"] == "r.patil@dbit.edu" and created["user"]["mustChangePassword"]
    assert created["temporaryPassword"] == "Temp-Rpatil-2345"
    assert (await s.post(f"{API}/admin/users", json={
        "email": "R.PATIL@dbit.edu", "name": "Again", "role": "faculty", "password": "Temp-Again-2345",
    })).status_code == 409
    assert (await s.post(f"{API}/admin/users", json={
        "email": "bad", "name": "X", "role": "faculty", "password": "Temp-Bad-2345",
    })).status_code == 400

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as person:
        body = (await person.post(f"{API}/session",
                                  json={"email": "r.patil@dbit.edu", "password": "Temp-Rpatil-2345"})).json()
        assert body["next"] == "change_password"
        r = await s.post(f"{API}/admin/users/{created['user']['id']}/reset-password",
                         json={"password": "Temp-Reset-2345"})
        assert r.status_code == 200 and r.json()["temporaryPassword"] == "Temp-Reset-2345"
        assert (await person.get(f"{API}/me")).status_code == 401
        body = (await person.post(f"{API}/session",
                                  json={"email": "r.patil@dbit.edu", "password": "Temp-Reset-2345"})).json()
        assert body["next"] == "change_password"

    r = await s.post(f"{API}/admin/users/{created['user']['id']}/reset-password", json={"resetMfa": True})
    assert r.status_code == 200 and r.json()["temporaryPassword"] != "Temp-Reset-2345"
    me_id = (await s.get(f"{API}/me")).json()["id"]
    assert (await s.post(f"{API}/admin/users/{me_id}/reset-password", json={})).status_code == 403


async def test_role_change_needs_reason_is_versioned_and_revokes_sessions(client):
    s = await _ossc(client)
    u = await make_user(FACULTY, "faculty")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as person:
        await sign_in(person, FACULTY)
        r = await s.patch(f"{API}/admin/users/{u.id}", json={"expectedVersion": 1, "role": "facility_manager"})
        assert r.status_code == 400 and "reason" in r.json()["field_errors"]
        r = await s.patch(f"{API}/admin/users/{u.id}", json={"expectedVersion": 1, "role": "facility_manager", "reason": "Promoted"})
        assert r.status_code == 200 and r.json()["version"] == 2
        # Old faculty session must not carry over into an admin role without TOTP.
        assert (await person.get(f"{API}/me")).status_code == 401
    stale = await s.patch(f"{API}/admin/users/{u.id}", json={"expectedVersion": 1, "name": "New Name"})
    assert stale.status_code == 409
    async with sessionmaker()() as db:
        ev = await db.scalar(select(AuditEvent).where(AuditEvent.action == "user.role_changed"))
        assert ev.before["role"] == "faculty" and ev.after["role"] == "facility_manager" and ev.reason == "Promoted"


async def test_email_change_revokes_sessions_and_moves_login(client):
    s = await _ossc(client)
    u = await make_user(FACULTY, "faculty")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as person:
        await sign_in(person, FACULTY)
        r = await s.patch(f"{API}/admin/users/{u.id}", json={"expectedVersion": 1, "email": "renamed@dbit.edu"})
        assert r.status_code == 200, r.text
        assert r.json()["email"] == "renamed@dbit.edu"
        assert (await person.get(f"{API}/me")).status_code == 401
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as person:
        assert (await person.post(f"{API}/session", json={"email": FACULTY, "password": PASSWORD})).status_code == 401
        assert (await person.post(f"{API}/session", json={"email": "renamed@dbit.edu", "password": PASSWORD})).status_code == 200


async def test_cannot_lock_out_last_admin_or_self(client):
    s = await _ossc(client)
    me = (await s.get(f"{API}/me")).json()
    r = await s.patch(f"{API}/admin/users/{me['id']}", json={"expectedVersion": 1, "active": False, "reason": "x"})
    assert r.status_code == 403
    other = await make_user("admin2@dbit.edu", "ossc")
    r = await s.patch(f"{API}/admin/users/{other.id}", json={"expectedVersion": 1, "active": False, "reason": "Left"})
    assert r.status_code == 200  # two admins existed


async def test_list_users_search_escapes_wildcards(client):
    s = await _ossc(client)
    await make_user("asha_k@dbit.edu", "faculty", "Asha_Kumar")
    await make_user("ashok@dbit.edu", "faculty", "Ashok Rao")
    r = await s.get(f"{API}/admin/users", params={"q": "asha_"})
    assert [u["email"] for u in r.json()["items"]] == ["asha_k@dbit.edu"]
    body = r.json()["items"][0]
    assert "passwordHash" not in body and "totpSecret" not in body and "tempPasswordEnc" not in body


async def test_credentials_export_requires_own_password_and_hides_chosen_password(client):
    s = await _ossc(client)
    temp = await make_user("temp@dbit.edu", "faculty", "Temp Person", password="Temp-Person-2345", temporary=True)
    await make_user("chosen@dbit.edu", "faculty", "Chosen Person")
    expired = await make_user("expired@dbit.edu", "faculty", "Expired Person", password="Temp-Expired-2345",
                              temporary=True)
    async with sessionmaker()() as db:
        (await db.get(User, expired.id)).temp_password_expires_at = now() - timedelta(minutes=1)
        await db.commit()

    assert (await s.post(f"{API}/admin/users/credentials/export", json={"password": "wrong"})).status_code == 400
    r = await s.post(f"{API}/admin/users/credentials/export", json={"password": PASSWORD})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(XLSX) and "no-store" in r.headers["cache-control"]
    assert r.headers["content-disposition"].startswith('attachment; filename="pcms-credentials-')
    assert r.headers["content-disposition"].endswith('.xlsx"')
    rows = {row["email"]: row for row in _rows(r)}
    assert rows["temp@dbit.edu"]["password"] == "Temp-Person-2345" and rows["temp@dbit.edu"]["status"] == "Not activated"
    assert rows["chosen@dbit.edu"]["password"] == "" and rows["chosen@dbit.edu"]["status"] == "Active"
    assert "Changed by employee" in rows["chosen@dbit.edu"]["notes"]
    assert rows["expired@dbit.edu"]["password"] == "" and rows["expired@dbit.edu"]["status"] == "Temporary password expired"
    assert "reset required" in rows["expired@dbit.edu"]["notes"]

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as person:
        body = (await person.post(f"{API}/session", json={"email": "temp@dbit.edu", "password": "Temp-Person-2345"})).json()
        r = await person.post(f"{API}/session/password", json={"currentPassword": "Temp-Person-2345", "newPassword": "Own-Choice-2345"},
                              headers={"X-CSRF-Token": body["csrfToken"]})
        assert r.status_code == 200

    r = await s.post(f"{API}/admin/users/credentials/export", json={"password": PASSWORD, "q": "temp@"})
    rows = _rows(r)
    assert len(rows) == 1 and rows[0]["email"] == "temp@dbit.edu" and rows[0]["password"] == ""

    async with sessionmaker()() as db:
        u = await db.get(User, temp.id)
        assert u.temp_password_enc is None and u.password_changed_at is not None
        ev = await db.scalar(select(AuditEvent).where(AuditEvent.action == "user.credentials_exported"))
        assert ev is not None and "Temp-Person-2345" not in str(ev.after)
